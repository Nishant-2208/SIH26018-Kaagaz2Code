"""
preprocessor.py — OpenCV preprocessing for scanned land-record images.

Pipeline: decode -> downscale -> grayscale -> denoise -> deskew -> CLAHE -> binarize.

Design note: we return BOTH the CLAHE grayscale and the binarized image. Hard
binarization helps on clean typed records but destroys thin Devanagari strokes on
faded/handwritten ones, where PaddleOCR does measurably better on the CLAHE image.
Which one is fed to OCR is an env switch (OCR_INPUT), not a code change.
"""

import io
import logging
import os
import re
from dataclasses import dataclass
from typing import Literal

import cv2
import numpy as np

logger = logging.getLogger(__name__)


def normalize_devanagari(text: str) -> str:
    """Deduplicate faux-bold repeating Devanagari syllables/clusters."""
    if not text:
        return ""
    # Matches any Devanagari syllable / cluster repeated immediately (e.g. ज़ि + ज़ि -> ज़ि, सं + सं -> सं)
    cluster = r"([\u0900-\u097f](?:[\u093c])?(?:[\u094d][\u0900-\u097f](?:[\u093c])?)*(?:[\u093e-\u094f\u0955-\u0957\u0962\u0963])?(?:[\u0901-\u0903])?)"
    return re.sub(rf"{cluster}\1+", r"\1", text)

# --- tunables -----------------------------------------------------------------
MAX_EDGE_PX: int = int(os.getenv("PREPROCESS_MAX_EDGE_PX", "2000"))
MAX_SKEW_DEG: float = float(os.getenv("PREPROCESS_MAX_SKEW_DEG", "15.0"))
BLUR_THRESHOLD: float = float(os.getenv("PREPROCESS_BLUR_THRESHOLD", "60.0"))
OCR_INPUT: Literal["clahe", "binary"] = os.getenv("OCR_INPUT", "clahe")  # type: ignore[assignment]


class ImageDecodeError(ValueError):
    """Raised when the uploaded bytes are not a decodable image or document."""


@dataclass
class PreprocessResult:
    """Everything downstream stages (and the debug endpoint) need."""

    ocr_input: np.ndarray  # 3-channel BGR — PaddleOCR expects HxWx3
    clahe_gray: np.ndarray  # single channel
    binary: np.ndarray  # single channel, 0/255
    width: int
    height: int
    skew_angle_deg: float
    blur_score: float  # variance of Laplacian; low == blurry
    is_likely_illegible: bool
    downscaled: bool
    direct_text: str = ""


def decode_document(raw: bytes, filename: str = "") -> tuple[np.ndarray, str]:
    """Decode raw upload bytes from an image, PDF, or Word DOCX document.

    Returns:
        tuple[np.ndarray, str]: (bgr_image, direct_text)
    Raises:
        ImageDecodeError: bytes are empty, corrupt, or an unsupported file.
    """
    if not raw:
        raise ImageDecodeError("Empty file body.")

    ext = os.path.splitext(filename)[1].lower() if filename else ""

    # 1. PDF detection (by extension or PDF magic header)
    if raw.startswith(b"%PDF") or ext == ".pdf":
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz  # type: ignore

            with fitz.open(stream=raw, filetype="pdf") as doc:
                if doc.page_count == 0:
                    raise ImageDecodeError("PDF contains no pages.")

                # Extract digital text if present and normalize faux-bold glyph duplication
                pdf_text_parts = []
                for page in doc:
                    t = page.get_text().strip()
                    if t:
                        pdf_text_parts.append(normalize_devanagari(t))
                pdf_text = "\n\n".join(pdf_text_parts)

                # Render page(s) to BGR image
                rendered = []
                for page in doc:
                    pix = page.get_pixmap(dpi=150)
                    im = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                    if pix.n == 4:
                        im = cv2.cvtColor(im, cv2.COLOR_RGBA2BGR)
                    elif pix.n == 3:
                        im = cv2.cvtColor(im, cv2.COLOR_RGB2BGR)
                    elif pix.n == 1:
                        im = cv2.cvtColor(im, cv2.COLOR_GRAY2BGR)
                    rendered.append(im)
                    if len(rendered) >= 5:  # cap at 5 pages
                        break

                if not rendered:
                    raise ImageDecodeError("Could not render any pages from PDF.")

                if len(rendered) == 1:
                    composite = rendered[0]
                else:
                    target_w = max(p.shape[1] for p in rendered)
                    standardized = []
                    for p in rendered:
                        if p.shape[1] != target_w:
                            scale = target_w / float(p.shape[1])
                            p = cv2.resize(p, (target_w, int(p.shape[0] * scale)), interpolation=cv2.INTER_AREA)
                        standardized.append(p)
                    composite = cv2.vconcat(standardized)

                direct = pdf_text if len(pdf_text) >= 80 else ""
                return composite, direct
        except Exception as exc:
            if isinstance(exc, ImageDecodeError):
                raise
            raise ImageDecodeError(f"Failed to decode PDF: {exc}") from exc

    # 2. DOCX detection (by extension or zip magic bytes PK\x03\x04)
    if ext in (".docx", ".doc") or (raw.startswith(b"PK\x03\x04") and ext != ".zip"):
        try:
            import docx

            doc = docx.Document(io.BytesIO(raw))
            lines = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    lines.append(t)
            for tbl in doc.tables:
                for row in tbl.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            direct_text = "\n".join(lines)

            # Check for embedded images in the DOCX
            embedded_imgs = []
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    try:
                        im = cv2.imdecode(np.frombuffer(rel.target_part.blob, dtype=np.uint8), cv2.IMREAD_COLOR)
                        if im is not None and im.shape[0] >= 32 and im.shape[1] >= 32:
                            embedded_imgs.append(im)
                    except Exception:
                        pass

            if embedded_imgs:
                target_w = max(p.shape[1] for p in embedded_imgs)
                standardized = []
                for p in embedded_imgs:
                    if p.shape[1] != target_w:
                        scale = target_w / float(p.shape[1])
                        p = cv2.resize(p, (target_w, int(p.shape[0] * scale)), interpolation=cv2.INTER_AREA)
                    standardized.append(p)
                composite = cv2.vconcat(standardized) if len(standardized) > 1 else standardized[0]
                return composite, direct_text

            if direct_text:
                canvas = np.full((720, 1000, 3), 255, dtype=np.uint8)
                return canvas, direct_text

            raise ImageDecodeError("DOCX contains neither readable text nor embedded images.")
        except Exception as exc:
            if isinstance(exc, ImageDecodeError):
                raise
            if ext in (".docx", ".doc"):
                raise ImageDecodeError(f"Failed to read DOCX file: {exc}") from exc

    # 3. Standard image (JPEG, PNG, WEBP, TIFF, BMP)
    img = decode_image(raw)
    return img, ""


def decode_image(raw: bytes) -> np.ndarray:
    """Decode raw upload bytes into a BGR ndarray.

    Raises:
        ImageDecodeError: bytes are empty, truncated, or not an image format
            OpenCV can read (guards against a renamed .pdf or a zero-byte upload).
    """
    if not raw:
        raise ImageDecodeError("Empty file body.")
    buf = np.frombuffer(raw, dtype=np.uint8)
    img = cv2.imdecode(buf, cv2.IMREAD_COLOR)
    if img is None:
        raise ImageDecodeError("Bytes could not be decoded as an image (corrupt or unsupported format).")
    if img.shape[0] < 32 or img.shape[1] < 32:
        raise ImageDecodeError(f"Image too small to contain text: {img.shape[1]}x{img.shape[0]}px.")
    return img


def _downscale(img: np.ndarray, max_edge: int = MAX_EDGE_PX) -> tuple[np.ndarray, bool]:
    """Cap the long edge. Phone photos arrive at 4000px+ and cost ~4x OCR time
    for no accuracy gain once text height is above ~25px."""
    h, w = img.shape[:2]
    longest = max(h, w)
    if longest <= max_edge:
        return img, False
    scale = max_edge / float(longest)
    resized = cv2.resize(img, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)
    logger.debug("Downscaled %dx%d -> %dx%d", w, h, resized.shape[1], resized.shape[0])
    return resized, True


def _estimate_skew(gray: np.ndarray) -> float:
    """Estimate page skew in degrees via the minimum-area rect of ink pixels.

    Returns 0.0 when the estimate is unreliable (too few ink pixels) or beyond
    MAX_SKEW_DEG — a 40-degree "correction" on a misread is far worse than none.
    """
    thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    coords = cv2.findNonZero(thresh)
    if coords is None or len(coords) < 100:
        logger.debug("Skew estimation skipped: insufficient ink pixels.")
        return 0.0

    angle = cv2.minAreaRect(coords)[-1]
    # OpenCV changed this API's range between versions: 4.5+ returns (0, 90],
    # older builds return [-90, 0). Normalising into (-45, 45] is correct for
    # both and avoids a spurious 90-degree "skew" on near-axis-aligned pages.
    while angle < -45.0:
        angle += 90.0
    while angle > 45.0:
        angle -= 90.0

    if abs(angle) > MAX_SKEW_DEG:
        logger.warning("Skew estimate %.2f deg exceeds clamp %.1f; not rotating.", angle, MAX_SKEW_DEG)
        return 0.0
    return float(angle)


def _rotate(img: np.ndarray, angle_deg: float) -> np.ndarray:
    """Rotate about the centre, expanding the canvas so no corner is clipped."""
    if abs(angle_deg) < 0.1:
        return img
    h, w = img.shape[:2]
    centre = (w / 2.0, h / 2.0)
    mat = cv2.getRotationMatrix2D(centre, angle_deg, 1.0)

    cos, sin = abs(mat[0, 0]), abs(mat[0, 1])
    new_w = int(h * sin + w * cos)
    new_h = int(h * cos + w * sin)
    mat[0, 2] += (new_w / 2.0) - centre[0]
    mat[1, 2] += (new_h / 2.0) - centre[1]

    border = 255 if img.ndim == 2 else (255, 255, 255)
    return cv2.warpAffine(
        img, mat, (new_w, new_h),
        flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_CONSTANT, borderValue=border,
    )


def preprocess(raw: bytes, filename: str = "") -> PreprocessResult:
    """Run the full preprocessing chain on raw upload bytes (image, PDF, or DOCX).

    Raises:
        ImageDecodeError: propagated from decode_document / decode_image.
    """
    img, direct_text = decode_document(raw, filename=filename)
    img, downscaled = _downscale(img)

    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)

    # Non-local-means: preserves stroke edges much better than a Gaussian blur,
    # which matters for the thin horizontal shirorekha in Devanagari script.
    denoised = cv2.fastNlMeansDenoising(gray, None, h=10, templateWindowSize=7, searchWindowSize=21)

    skew = _estimate_skew(denoised)
    deskewed = _rotate(denoised, skew)

    # CLAHE instead of global equalisation: scans are unevenly lit (book spine
    # shadow, flash hotspot) and global EQ blows out the bright half.
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(deskewed)

    binary = cv2.adaptiveThreshold(
        clahe, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, blockSize=31, C=11
    )

    blur_score = float(cv2.Laplacian(clahe, cv2.CV_64F).var())
    # Digital text documents (like DOCX or digital PDFs) should not be flagged as illegible
    illegible = blur_score < BLUR_THRESHOLD and not bool(direct_text)

    chosen = binary if OCR_INPUT == "binary" else clahe
    ocr_input = cv2.cvtColor(chosen, cv2.COLOR_GRAY2BGR)

    h, w = clahe.shape[:2]
    logger.info(
        "Preprocessed: %dx%d skew=%.2fdeg blur=%.1f illegible=%s input=%s direct_text_len=%d",
        w, h, skew, blur_score, illegible, OCR_INPUT, len(direct_text),
    )

    return PreprocessResult(
        ocr_input=ocr_input,
        clahe_gray=clahe,
        binary=binary,
        width=w,
        height=h,
        skew_angle_deg=round(skew, 2),
        blur_score=round(blur_score, 2),
        is_likely_illegible=illegible,
        downscaled=downscaled,
        direct_text=direct_text,
    )
